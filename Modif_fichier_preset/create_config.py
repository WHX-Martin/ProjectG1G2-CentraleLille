from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter


def create_excel_template(output_path):
    """create default config Excel"""
    wb = Workbook()

    # ==================== create Sensors table ====================
    sensors_ws = wb.active
    sensors_ws.title = "Sensors"

    # row 1
    sensors_ws.cell(row=1, column=2, value="Position(m)").alignment = Alignment(
        horizontal="center"
    )
    sensors_ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
    # row 2
    headers = ["Name", "x", "y", "z"]
    for col, header in enumerate(headers, start=1):
        sensors_ws.cell(row=2, column=col, value=header)

    # set width
    for col in range(1, 5):
        sensors_ws.column_dimensions[get_column_letter(col)].width = 12

    # ==================== create Spotlights table ====================
    spotlights_ws = wb.create_sheet("Spotlights")

    # row 1

    spotlights_ws.cell(row=1, column=2, value="Theoretical Position(m)")
    spotlights_ws.cell(row=1, column=5, value="Axis(pan0,tilt90)")
    spotlights_ws.cell(row=1, column=6, value="Rotation")
    spotlights_ws.cell(row=1, column=7, value="Orientation pan")
    spotlights_ws.cell(row=1, column=8, value="0:anticlockwise")
    spotlights_ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)

    # row 2
    headers = [
        "ID",
        "x",
        "y",
        "z",
        "choice:x/y/-x/-y",
        "",
        "",
        "1:clockwise",
    ]
    for col, header in enumerate(headers, start=1):
        spotlights_ws.cell(row=2, column=col, value=header)

    # set alignment
    for row in spotlights_ws.iter_rows(min_row=1, max_row=2):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")

    # set width
    col_widths = [8, 12, 12, 12, 18, 15, 17, 17]
    for idx, width in enumerate(col_widths, start=1):
        spotlights_ws.column_dimensions[get_column_letter(idx)].width = width

    spotlights_ws.merge_cells(start_row=1, start_column=6, end_row=2, end_column=6)
    spotlights_ws.merge_cells(start_row=1, start_column=7, end_row=2, end_column=7)
    # save
    wb.save(output_path)
    print(f"Successfully Create ：{output_path}")


if __name__ == "__main__":
    create_excel_template("config-3.0.0-auto.xlsx")
