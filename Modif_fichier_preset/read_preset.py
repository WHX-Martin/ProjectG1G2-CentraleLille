import xml.etree.ElementTree as ET
from openpyxl import Workbook, load_workbook
import os
import glob

def readXml(path):
    dictXml = {}
    file_name = os.path.basename(path)
    # Parsing an XML file
    tree = ET.parse(path)  # XMLpath
    root = tree.getroot()
    presetName=file_name
    print(presetName)
    dictXml['Name'] = presetName
    # Traverse each Phaser element.
    for phaser in root.findall('.//Phaser'):
        ID = phaser.get('ID')
        attribute = phaser.get('Attribute')
        
        # Find the Step element and retrieve the AbsolutePhys value.
        step = phaser.find('Step')
        absolute_phys = step.get('AbsolutePhys')
        
        # Output the ID, Attribute, and AbsolutePhys values.
        print("ID:", ID)
        print("Attribute:", attribute)
        print("AbsolutePhys:", absolute_phys)
        if str(ID) not in dictXml.keys():
            dictXml[str(ID)]={}
        valeur = dictXml[str(ID)]
        if attribute not in valeur.keys():
            valeur[attribute] = absolute_phys
            dictXml[str(ID)] = valeur
    
    print(dictXml)
    return dictXml


def writeExcel(data,excel_path):
    if excel_path == None:
        excel_path='Presets.xlsx'

    # Check if the file exists.
    if os.path.exists(excel_path):
        # Open existing Excel file 
        wb = load_workbook(excel_path)
    else:
        # create a new Workbook
        wb = Workbook()
        default_sheet = wb.sheetnames[0]
        wb.remove(wb[default_sheet])
    row_num = 1
    # Iterate through the data and create different tables based on the Name.
    for key, value in data.items():
        if key == 'Name':
            name = value
            # Check if there is a table named "name".
            if name not in wb.sheetnames:
                # create a new table
                ws = wb.create_sheet(title=name)
                # write the table head
                ws['A1'] = 'ID'
                ws['B1'] = 'Pan'
                ws['C1'] = 'Tilt'
            else:
                # if the table exist, load the table
                ws = wb[name]
        else:
            # write the data
            pan = value['Pan']
            tilt = value['Tilt']
            row_num = row_num + 1
            ws.cell(row=row_num, column=1).value = key
            ws.cell(row=row_num, column=2).value = pan
            ws.cell(row=row_num, column=3).value = tilt

    # save the Excel
    wb.save(excel_path)


def get_xml_files(folder_path):
    xml_files = []
    # Use the glob module to retrieve all XML files
    for file_path in glob.glob(os.path.join(folder_path, '*.xml')):
        xml_files.append(file_path)
    return xml_files




#test
# Specify the folder path.
folder_path = '.'

# Get all the xml files in the folder 
xml_files = get_xml_files(folder_path)
print("xml file in the folder：", xml_files)
for file in xml_files:
    data = readXml(file)
    writeExcel(data,'Presets1.xlsx')

